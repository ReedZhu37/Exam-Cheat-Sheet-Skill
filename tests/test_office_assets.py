import hashlib
from pathlib import Path

import pytest
from PIL import Image
from docx import Document
from docx.shared import Inches as WordInches
from pptx import Presentation
from pptx.util import Inches

from exam_cheat_sheet.ingest import ingest_source


@pytest.mark.parametrize('extension', ['docx', 'pptx'])
def test_embedded_diagram_is_preserved_with_digest_and_review_flag(tmp_path, extension):
    image = tmp_path / 'diagram.png'
    Image.new('RGB', (60, 40), 'blue').save(image)
    document = tmp_path / f'notes.{extension}'
    if extension == 'docx':
        doc = Document()
        doc.add_paragraph('Question: interpret the attached blue diagram.')
        doc.add_picture(str(image), width=WordInches(1))
        doc.save(document)
    else:
        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[6])
        slide.shapes.add_picture(str(image), Inches(1), Inches(1))
        doc.save(document)
    result = ingest_source({
        'id': 'notes', 'path': document.name, 'role': 'notes', 'split': 'train',
        'status': 'pending', 'metadata': None,
        'provenance': {'origin': 'synthetic', 'license': 'CC0-1.0', 'acquired_at': None},
    }, tmp_path, tmp_path / 'extracted')
    page = result['metadata']['pages'][0]
    assets = [b for b in page['blocks'] if b.get('type') == 'image']
    assert len(assets) == 1
    assert Path(assets[0]['asset_path']).read_bytes() == image.read_bytes()
    assert assets[0]['sha256'] == hashlib.sha256(image.read_bytes()).hexdigest()
    assert page['quality'] == 'needs_review'
