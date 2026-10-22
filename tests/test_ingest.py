import copy
import hashlib
from pathlib import Path

import pytest
from PIL import Image
from reportlab.pdfgen import canvas

from exam_cheat_sheet import IngestError, ManifestError, ingest_source, validate_manifest


def entry(path, role="notes", split="train"):
    return {
        "id": "synthetic", "path": path, "role": role, "split": split,
        "status": "pending", "metadata": None,
        "provenance": {"origin": "Locally generated synthetic fixture", "license": "CC0-1.0", "acquired_at": None},
    }


def make_pdf(path, texts, *, figure=False):
    writer = canvas.Canvas(str(path), pagesize=(300, 400))
    for text in texts:
        if text:
            writer.drawString(30, 350, text)
        if figure:
            writer.circle(120, 150, 60, fill=1)
        writer.showPage()
    writer.save()


def test_pdf_preserves_every_page_and_full_page_raster(tmp_path):
    path = tmp_path / "synthetic.pdf"
    make_pdf(path, ["First page", "Second page", "Third page"])
    source = entry(path.name)
    result = ingest_source(source, tmp_path, tmp_path / "out")
    metadata = result["metadata"]
    assert source["metadata"] is None
    assert result["status"] == "validated"
    assert result["provenance"]["acquired_at"]
    assert metadata["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert metadata["size_bytes"] == path.stat().st_size
    assert metadata["page_count"] == 3
    assert metadata["page_unit"] == "pdf_page"
    assert "pypdfium2=" in metadata["parser"]["version"]
    assert [p["anchor"] for p in metadata["pages"]] == ["page:1", "page:2", "page:3"]
    for page in metadata["pages"]:
        assert "page" in page["text"]
        asset = Path(page["raster_path"])
        assert asset.is_file()
        with Image.open(asset) as image:
            assert image.size == (600, 800)
        assert page["page_sha256"] == hashlib.sha256(asset.read_bytes()).hexdigest()
        assert page["quality"] == "needs_review"


def test_missing_pdf_ledger_page_rejected(tmp_path):
    make_pdf(tmp_path / "synthetic.pdf", ["one", "two", "three"])
    result = ingest_source(entry("synthetic.pdf"), tmp_path, tmp_path / "out")
    result["metadata"]["pages"].pop(1)
    with pytest.raises(ManifestError, match="Missing"):
        validate_manifest({"schema_version": "1.0", "sources": [result]})


def test_duplicate_pages_detected_and_retained(tmp_path):
    make_pdf(tmp_path / "synthetic.pdf", ["repeated", "repeated", "unique"])
    result = ingest_source(entry("synthetic.pdf"), tmp_path, tmp_path / "out")
    pages = result["metadata"]["pages"]
    assert len(pages) == 3
    assert [page["duplicate_of"] for page in pages] == [None, 1, None]
    assert any("Duplicate" in reason for reason in pages[1]["review_reasons"])
    broken = copy.deepcopy(result)
    broken["metadata"]["pages"][1]["duplicate_of"] = None
    with pytest.raises(ManifestError, match="Duplicate"):
        validate_manifest({"schema_version": "1.0", "sources": [broken]})


@pytest.mark.parametrize("text,figure", [("E = mc^2", False), ("", True)])
def test_math_and_figure_pages_need_review_without_ocr(tmp_path, text, figure):
    make_pdf(tmp_path / "synthetic.pdf", [text], figure=figure)
    result = ingest_source(entry("synthetic.pdf"), tmp_path, tmp_path / "out")
    page = result["metadata"]["pages"][0]
    assert page["quality"] == "needs_review"
    assert any("no image OCR performed" in reason for reason in page["review_reasons"])
    if not text:
        assert not page["text"].strip()
        assert any("No embedded text" in reason for reason in page["review_reasons"])
    assert Path(page["raster_path"]).is_file()


def test_corrupt_pdf_has_explicit_failure_without_published_assets(tmp_path):
    (tmp_path / "broken.pdf").write_bytes(b"%PDF-1.7\nnot a real PDF")
    out = tmp_path / "out"
    with pytest.raises(IngestError, match="Cannot parse"):
        ingest_source(entry("broken.pdf"), tmp_path, out)
    assert not list(out.iterdir())


def test_denied_material_never_opened(tmp_path, monkeypatch):
    path = tmp_path / "hidden.txt"
    path.write_text("synthetic holdout", encoding="utf-8")
    original = Path.open

    def guarded_open(self, *args, **kwargs):
        if self == path:
            pytest.fail("Denied source content was opened")
        return original(self, *args, **kwargs)

    monkeypatch.setattr(Path, "open", guarded_open)
    with pytest.raises(ManifestError, match="cannot read"):
        ingest_source(entry(path.name, "past_exam", "sealed_test"), tmp_path, tmp_path / "out")
    assert not (tmp_path / "out").exists()


@pytest.mark.parametrize("role,path", [
    ("solution", "train/solutions/archival.txt"),
    ("past_exam", "train/questions/archival.txt"),
    ("notes", "train/notes.txt"),
])
def test_builder_train_material_needs_no_freeze_or_dev_gate(tmp_path, monkeypatch, role, path):
    from exam_cheat_sheet import compile as compiler

    def forbidden_verification(*args, **kwargs):
        pytest.fail("Train ingestion must not invoke freeze verification")

    monkeypatch.setattr(compiler, "verify_freeze", forbidden_verification)
    material = tmp_path / path
    material.parent.mkdir(parents=True, exist_ok=True)
    material.write_text("Synthetic archival learning material", encoding="utf-8")
    result = ingest_source(entry(path, role), tmp_path, tmp_path / "out")
    assert result["status"] == "validated"
    assert result["role"] == role
    assert result["split"] == "train"
    assert result["metadata"]["pages"][0]["text"] == "Synthetic archival learning material"


@pytest.fixture
def synthetic_freeze(tmp_path):
    from exam_cheat_sheet.compile import freeze_sheet

    sheet = tmp_path / "sheet"
    sheet.mkdir()
    (sheet / "cheatsheet.md").write_text("Synthetic frozen sheet", encoding="utf-8")
    frozen = tmp_path / "freeze.json"
    freeze_sheet(sheet, [], frozen)
    return frozen


@pytest.mark.parametrize("access,role,path", [
    ("final_evaluator", "past_exam", "sealed_test/questions.txt"),
    ("final_grader", "past_exam", "sealed_test/questions.txt"),
    ("final_grader", "solution", "sealed_test/solutions/answer.txt"),
])
def test_final_access_reads_synthetic_material_after_real_freeze_verification(
        tmp_path, synthetic_freeze, access, role, path):
    material = tmp_path / path
    material.parent.mkdir(parents=True, exist_ok=True)
    material.write_text("Synthetic final material", encoding="utf-8")
    result = ingest_source(entry(path, role, "sealed_test"), tmp_path, tmp_path / "out",
                           access=access, freeze=synthetic_freeze)
    assert result["metadata"]["pages"][0]["text"] == "Synthetic final material"
    assert result["role"] == role
    assert result["split"] == "sealed_test"


@pytest.mark.parametrize("failure", ["absent", "missing", "malformed", "drift"])
@pytest.mark.parametrize("access", ["final_evaluator", "final_grader"])
def test_final_freeze_failure_precedes_any_material_resolution(
        tmp_path, synthetic_freeze, monkeypatch, failure, access):
    import exam_cheat_sheet.manifest as module

    frozen = synthetic_freeze
    if failure == "absent":
        frozen = None
    elif failure == "missing":
        frozen = tmp_path / "missing-freeze.json"
    elif failure == "malformed":
        frozen.write_text("not JSON", encoding="utf-8")
    else:
        (tmp_path / "sheet" / "cheatsheet.md").write_text("Changed after freezing", encoding="utf-8")

    def forbidden_resolution(*args, **kwargs):
        pytest.fail("Material path resolved before successful freeze verification")

    monkeypatch.setattr(module, "resolve_source", forbidden_resolution)
    with pytest.raises(ManifestError, match="freeze|Freeze"):
        ingest_source(entry("sealed_test/unread.txt", "past_exam", "sealed_test"),
                      tmp_path, tmp_path / "out", access=access, freeze=frozen)
    assert not (tmp_path / "out").exists()


@pytest.mark.parametrize("access,role,split", [
    ("builder", "past_exam", "sealed_test"),
    ("builder", "solution", "dev"),
    ("evaluator", "solution", "dev"),
    ("evaluator", "past_exam", "sealed_test"),
    ("grader", "solution", "sealed_test"),
    ("final_evaluator", "solution", "sealed_test"),
])
def test_valid_freeze_cannot_override_role_policy(
        tmp_path, synthetic_freeze, monkeypatch, access, role, split):
    import exam_cheat_sheet.manifest as module

    def forbidden_resolution(*args, **kwargs):
        pytest.fail("Disallowed material path was resolved")

    monkeypatch.setattr(module, "resolve_source", forbidden_resolution)
    with pytest.raises(ManifestError, match="cannot read"):
        ingest_source(entry("unread.txt", role, split), tmp_path, tmp_path / "out",
                      access=access, freeze=synthetic_freeze)


def test_docx_paragraphs_tables_and_equations(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement

    document = Document()
    paragraph = document.add_paragraph("A synthetic paragraph")
    paragraph._p.append(OxmlElement("m:oMath"))
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    document.add_paragraph("After table")
    document.save(tmp_path / "synthetic.docx")
    result = ingest_source(entry("synthetic.docx"), tmp_path, tmp_path / "out")
    metadata = result["metadata"]
    page = metadata["pages"][0]
    assert metadata["page_unit"] == "document"
    assert metadata["page_count"] == 1
    assert [b["type"] for b in page["blocks"]] == ["paragraph", "table", "paragraph"]
    assert page["blocks"][1]["rows"][0][1]["anchor"] == "table:1/row:1/cell:2"
    assert "left\tright" in page["text"]
    assert any("Office math" in reason for reason in page["review_reasons"])


def test_pptx_slide_and_shape_anchors(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    for text in ("Synthetic slide one", "Synthetic slide two"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text = text
    presentation.save(tmp_path / "synthetic.pptx")
    result = ingest_source(entry("synthetic.pptx", "slides"), tmp_path, tmp_path / "out")
    metadata = result["metadata"]
    assert metadata["page_count"] == 2
    assert metadata["page_unit"] == "slide"
    for number, page in enumerate(metadata["pages"], 1):
        assert page["anchor"] == f"slide:{number}"
        assert page["blocks"][0]["anchor"].startswith(f"slide:{number}/shape:")
        assert page["quality"] == "needs_review"


@pytest.mark.parametrize("extension", ["md", "txt"])
def test_text_source_and_changed_acquisition_rejected(tmp_path, extension):
    path = tmp_path / f"synthetic.{extension}"
    path.write_text("# Topic\nA = B\n", encoding="utf-8")
    result = ingest_source(entry(path.name), tmp_path, tmp_path / "out")
    page = result["metadata"]["pages"][0]
    assert page["blocks"][1]["anchor"] == "line:2"
    assert page["quality"] == "needs_review"
    path.write_text("changed", encoding="utf-8")
    with pytest.raises(IngestError, match="checksum"):
        ingest_source(result, tmp_path, tmp_path / "out")


def test_unsupported_format_fails_explicitly(tmp_path):
    (tmp_path / "unsupported.html").write_text("synthetic", encoding="utf-8")
    with pytest.raises(IngestError, match="Unsupported"):
        ingest_source(entry("unsupported.html"), tmp_path, tmp_path / "out")
