"""Local ingestion; no network acquisition and no image OCR.

PDF pages retain full-page PNG renders and always need human review for formulas,
figures, reading order and missing text. DOCX and plain text have one *logical
document* ledger entry, never an invented physical page count. PPTX uses slides.
Returned dictionaries are source entries suitable for a validated manifest.
"""

from __future__ import annotations

import copy
import hashlib
import io
import os
import re
import tempfile
from datetime import datetime, timezone
from importlib.metadata import version
from pathlib import Path

from .manifest import ManifestError, authorize_source, resolve_source, validate_manifest


class IngestError(ValueError):
    """Material cannot be faithfully parsed or its identity has changed."""


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _page(number: int, text: str, anchor: str, *, reasons: list[str] | None = None,
          raster: str | None = None, digest: str | None = None,
          blocks: list[dict] | None = None) -> dict:
    reasons = list(reasons or [])
    if re.search(r"[$=∑∫√]|\\(?:frac|sum|int|sqrt)\b", text):
        reasons.append("Formula notation requires fidelity review")
    result = {
        "page_number": number, "anchor": anchor, "text": text,
        "text_sha256": _sha(text.encode("utf-8")),
        "page_sha256": digest or _sha(text.encode("utf-8")),
        "raster_path": raster,
        "quality": "needs_review" if reasons else "extracted",
        "review_reasons": reasons, "duplicate_of": None,
    }
    if blocks is not None:
        result["blocks"] = blocks
    return result


def _pdf(blob: bytes, staging: Path, destination: Path) -> tuple[list[dict], str, str, str]:
    from pypdf import PdfReader
    import pypdfium2 as pdfium

    reader = PdfReader(io.BytesIO(blob), strict=True)
    if reader.is_encrypted and not reader.decrypt(""):
        raise IngestError("Password-protected PDF is unsupported")
    pages = []
    with pdfium.PdfDocument(blob) as rendered:
        if len(reader.pages) != len(rendered) or not len(rendered):
            raise IngestError("PDF parser/render page counts disagree or PDF is empty")
        for index, original in enumerate(reader.pages):
            text = original.extract_text() or ""
            name = f"page-{index + 1:04d}.png"
            page = rendered[index]
            try:
                bitmap = page.render(scale=2, draw_annots=True)
                try:
                    image = bitmap.to_pil()
                    try:
                        image.save(staging / name, format="PNG")
                    finally:
                        image.close()
                finally:
                    bitmap.close()
            finally:
                page.close()
            reasons = ["Full-page visual review required for math, figures and reading order; no image OCR performed"]
            if not text.strip():
                reasons.append("No embedded text extracted; inspect raster or arrange separate OCR")
            pages.append(_page(index + 1, text, f"page:{index + 1}", reasons=reasons,
                               raster=str(destination / name), digest=_sha((staging / name).read_bytes())))
    return pages, "pdf_page", "pypdf+pypdfium2", f"pypdf={version('pypdf')};pypdfium2={version('pypdfium2')}"


def _docx(blob: bytes, staging: Path, destination: Path) -> tuple[list[dict], str, str, str]:
    from docx import Document
    from docx.table import Table

    document = Document(io.BytesIO(blob))
    blocks, texts = [], []
    paragraph_number = table_number = 0
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            table_number += 1
            rows = []
            for row_number, row in enumerate(item.rows, 1):
                cells = []
                for column_number, cell in enumerate(row.cells, 1):
                    cells.append({"anchor": f"table:{table_number}/row:{row_number}/cell:{column_number}", "text": cell.text})
                rows.append(cells)
                texts.append("\t".join(cell["text"] for cell in cells))
            blocks.append({"type": "table", "anchor": f"table:{table_number}", "rows": rows})
        else:
            paragraph_number += 1
            blocks.append({"type": "paragraph", "anchor": f"paragraph:{paragraph_number}", "text": item.text})
            texts.append(item.text)
    reasons = ["Word layout and physical pagination are not available; review document formatting and visual content"]
    if document.element.xpath("//*[local-name()='oMath' or local-name()='oMathPara']"):
        reasons.append("Office math found; equation XML is not converted to plain text")
    if document.element.xpath("//*[local-name()='drawing' or local-name()='pict']"):
        reasons.append("Embedded figures require visual review; no image OCR performed")
    for index, relation in enumerate(document.part.rels.values()):
        if not relation.is_external and relation.reltype.endswith('/image'):
            part = relation.target_part
            suffix = Path(str(part.partname)).suffix.lower()
            name = f'image-{index:04d}{suffix}'
            (staging / name).write_bytes(part.blob)
            blocks.append({'type': 'image', 'anchor': f'image:{relation.rId}',
                           'asset_path': str(destination / name), 'sha256': _sha(part.blob),
                           'content_type': part.content_type})
    page = _page(1, "\n".join(texts), "document:1", reasons=reasons, blocks=blocks)
    return [page], "document", "python-docx", version("python-docx")


def _pptx(blob: bytes, staging: Path, destination: Path) -> tuple[list[dict], str, str, str]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(io.BytesIO(blob))
    pages = []
    for number, slide in enumerate(presentation.slides, 1):
        blocks, texts = [], []

        def shapes(items, prefix):
            for shape in items:
                anchor = f"{prefix}/shape:{shape.shape_id}"
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    shapes(shape.shapes, anchor)
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    name = f'slide-{number:04d}-image-{shape.shape_id}.{shape.image.ext}'
                    (staging / name).write_bytes(shape.image.blob)
                    blocks.append({'type': 'image', 'anchor': anchor,
                                   'asset_path': str(destination / name), 'sha256': _sha(shape.image.blob),
                                   'content_type': shape.image.content_type})
                elif shape.has_table:
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    blocks.append({"type": "table", "anchor": anchor, "rows": rows})
                    texts.extend("\t".join(row) for row in rows)
                elif shape.has_text_frame:
                    blocks.append({"type": "text", "anchor": anchor, "text": shape.text})
                    texts.append(shape.text)
                else:
                    blocks.append({"type": "visual", "anchor": anchor})

        shapes(slide.shapes, f"slide:{number}")
        reasons = ["Slide figures, formulas and spatial reading order require visual review; no image OCR performed"]
        # Include relationship targets' bytes so slides with different pictures
        # but identical text are not falsely identified as duplicate pages.
        identity = hashlib.sha256(slide.part.blob)
        for relationship in sorted(slide.part.rels.values(), key=lambda rel: rel.rId):
            if not relationship.is_external:
                identity.update(relationship.target_part.blob)
        pages.append(_page(number, "\n".join(texts), f"slide:{number}", reasons=reasons,
                           digest=identity.hexdigest(), blocks=blocks))
    if not pages:
        raise IngestError("Presentation contains no slides")
    return pages, "slide", "python-pptx", version("python-pptx")


def _read_authorized(source: dict, root: Path, path: Path) -> tuple[Path, bytes]:
    # `path` is produced only after authorize_source has passed all access gates.
    if resolve_source(source, root) != path:
        raise ManifestError("Source path changed during authorization")
    before = path.stat()
    # Check again after stat so a changed ancestor cannot authorize an escape.
    if resolve_source(source, root) != path:
        raise ManifestError("Source path changed during authorization")
    with path.open("rb") as stream:
        opened = os.fstat(stream.fileno())
        if (before.st_dev, before.st_ino) != (opened.st_dev, opened.st_ino):
            raise ManifestError("Source identity changed before reading")
        blob = stream.read()
        after = os.fstat(stream.fileno())
    if (opened.st_size, opened.st_mtime_ns) != (after.st_size, after.st_mtime_ns) or len(blob) != after.st_size:
        raise IngestError("Source changed while reading")
    return path, blob


def ingest_source(source: dict, root: Path, out: Path, access: str = "builder", *,
                  freeze: Path | None = None) -> dict:
    """Authorize, snapshot and ingest one trusted manifest source entry.

    Validate the *whole* manifest with root first to detect contradictory aliases.
    Each entry requires id/path/role/split/status/provenance/metadata. Pending
    metadata and provenance.acquired_at may be null. Returns a deep-copied entry
    with status='validated', acquisition time, checksums, parser version and page
    ledger at metadata.pages. Validation means structural acquisition integrity,
    not human approval: quality flags remain needs_review where appropriate.
    PDF PNG paths are absolute under out; no output is published on parse failure.
    Existing acquisition metadata binds the bytes: checksum changes are rejected.
    Builder accepts explicitly manifested train solutions for learning; these
    remain role='solution', not independent question-frequency evidence. Final
    identities require freeze=Path(...), verified before source-path resolution
    or reads. Builder/train and ordinary dev access do not require a freeze.
    """
    path = authorize_source(source, root, access, freeze=freeze)
    validate_manifest({"schema_version": "1.0", "sources": [source]}, root=root)
    extension = path.suffix.casefold()
    if extension not in {".pdf", ".docx", ".pptx", ".md", ".txt"}:
        raise IngestError(f"Unsupported material format: {extension or '(none)'}")
    path, blob = _read_authorized(source, root, path)
    digest = _sha(blob)
    if source["metadata"] is not None:
        old = source["metadata"]
        if old["sha256"] != digest or old["size_bytes"] != len(blob):
            raise IngestError("Acquired source checksum or size changed")
    out = Path(out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    destination = out / f"{_sha(source['id'].encode())[:16]}-{digest}"
    with tempfile.TemporaryDirectory(prefix=".ingest-", dir=out) as temporary:
        staging = Path(temporary)
        try:
            if extension == ".pdf":
                pages, unit, parser, parser_version = _pdf(blob, staging, destination)
            elif extension == ".docx":
                pages, unit, parser, parser_version = _docx(blob, staging, destination)
            elif extension == ".pptx":
                pages, unit, parser, parser_version = _pptx(blob, staging, destination)
            else:
                text = blob.decode("utf-8-sig")
                blocks = [{"anchor": f"line:{n}", "text": line} for n, line in enumerate(text.splitlines(), 1)]
                pages = [_page(1, text, "document:1", blocks=blocks)]
                unit, parser, parser_version = "document", "utf-8-text", "1"
        except IngestError:
            raise
        except Exception as exc:
            raise IngestError(f"Cannot parse {source['path']}: {type(exc).__name__}: {exc}") from exc
        seen = {}
        for page in pages:
            previous = seen.setdefault(page["page_sha256"], page["page_number"])
            if previous != page["page_number"]:
                page["duplicate_of"] = previous
                page["quality"] = "needs_review"
                page["review_reasons"].append(f"Duplicate page content: matches page {previous}")
        result = copy.deepcopy(source)
        result["status"] = "validated"
        result["provenance"]["acquired_at"] = result["provenance"]["acquired_at"] or datetime.now(timezone.utc).isoformat()
        result["metadata"] = {
            "sha256": digest, "size_bytes": len(blob),
            "parser": {"name": parser, "version": parser_version},
            "page_count": len(pages), "page_unit": unit, "pages": pages,
        }
        validate_manifest({"schema_version": "1.0", "sources": [result]}, root=root)
        if extension in {".pdf", ".docx", ".pptx"}:
            destination.mkdir(exist_ok=True)
            for asset in staging.iterdir():
                os.replace(asset, destination / asset.name)
    return result
